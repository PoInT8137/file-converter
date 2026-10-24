import os
import subprocess
import sys
import tempfile
import wave
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
os.environ["FILECONVERTER_DATA"] = tempfile.mkdtemp(prefix="fc_test_data_")

DATA = ROOT / "tests" / "data"
# пробелы и кириллица в пути специально
OUT = ROOT / "tests" / "out" / "результаты тестов"


def make_pdf(path, pages=2):
    import pymupdf
    font = ROOT / "assets" / "fonts" / "IBMPlexMono-Regular.ttf"
    doc = pymupdf.open()
    for i in range(pages):
        page = doc.new_page()
        page.insert_font(fontname="plex", fontfile=str(font))
        page.insert_text((72, 100), f"Страница {i + 1}. Проверка конвертера File Converter.",
                         fontname="plex", fontsize=14)
        page.insert_text((72, 130), "Second line in English 12345", fontname="plex", fontsize=12)
        if i == 0:
            y = 200
            for r in range(4):
                for c in range(3):
                    rect = pymupdf.Rect(72 + c * 120, y + r * 24, 72 + (c + 1) * 120, y + (r + 1) * 24)
                    page.draw_rect(rect, color=(0, 0, 0), width=1)
                    page.insert_text((rect.x0 + 6, rect.y1 - 7), f"R{r}C{c}", fontname="plex", fontsize=10)
    doc.save(str(path))
    doc.close()


def make_docx(path):
    from docx import Document
    d = Document()
    d.add_heading("Тестовый документ", 1)
    d.add_paragraph("Обычный абзац текста для проверки конвертации. Second paragraph in English.")
    d.save(str(path))


def make_images(folder):
    from PIL import Image, ImageDraw
    img = Image.new("RGBA", (320, 200), (20, 24, 29, 255))
    d = ImageDraw.Draw(img)
    d.rectangle([20, 20, 300, 180], outline=(245, 165, 36, 255), width=6)
    d.text((40, 90), "File Converter test", fill=(230, 234, 240, 255))
    img.save(folder / "картинка.png")
    img.convert("RGB").save(folder / "фото.jpg", quality=90)


def make_wav(path, seconds=2):
    import math
    import struct
    rate = 22050
    with wave.open(str(path), "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(rate)
        frames = b"".join(struct.pack("<h", int(12000 * math.sin(2 * math.pi * 440 * i / rate)))
                          for i in range(rate * seconds))
        w.writeframes(frames)


def make_mp4(path):
    ff = ROOT / "engines" / "ffmpeg.exe"
    if not ff.exists():
        return
    subprocess.run([str(ff), "-y", "-hide_banner", "-loglevel", "error",
                    "-f", "lavfi", "-i", "testsrc=duration=2:size=320x240:rate=15",
                    "-f", "lavfi", "-i", "sine=frequency=440:duration=2",
                    "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-shortest", str(path)],
                   check=True, creationflags=0x08000000)


@pytest.fixture(scope="session")
def data_dir():
    folder = DATA / "тестовые файлы"
    folder.mkdir(parents=True, exist_ok=True)
    if not (folder / "документ.pdf").exists():
        make_pdf(folder / "документ.pdf")
    if not (folder / "документ.docx").exists():
        make_docx(folder / "документ.docx")
    if not (folder / "картинка.png").exists():
        make_images(folder)
    if not (folder / "звук.wav").exists():
        make_wav(folder / "звук.wav")
    if not (folder / "видео.mp4").exists():
        make_mp4(folder / "видео.mp4")
    return folder


@pytest.fixture(scope="session")
def out_dir():
    import shutil
    shutil.rmtree(OUT, ignore_errors=True)
    OUT.mkdir(parents=True, exist_ok=True)
    return OUT


@pytest.fixture(scope="session")
def engines():
    from fileconverter.config import Settings
    from fileconverter.engines.finder import Engines
    return Engines(Settings())


@pytest.fixture(scope="session")
def shared_app(tmp_path_factory):
    # одно окно на все проверки интерфейса, как в настоящей программе: несколько Tk подряд
    # в одном процессе на Windows иногда не могут прочитать tk.tcl
    os.environ["FILECONVERTER_DATA"] = str(tmp_path_factory.mktemp("ui_data"))
    from fileconverter.app import App
    a = App(selftest=True)
    a.update()
    yield a
    a.on_close()