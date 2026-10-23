import zipfile
from pathlib import Path

import pytest

from fileconverter.engines.base import Context
from fileconverter.naming import Namer
from fileconverter.registry import registry

# (исходный файл, формат результата)
CASES = [
    ("документ.pdf", "docx"), ("документ.pdf", "txt"), ("документ.pdf", "html"),
    ("документ.pdf", "png"), ("документ.pdf", "jpg"), ("документ.pdf", "pptx"), ("документ.pdf", "xlsx"),
    ("документ.docx", "pdf"), ("документ.docx", "odt"), ("документ.docx", "rtf"),
    ("документ.docx", "txt"), ("документ.docx", "html"), ("документ.docx", "md"),
    ("документ.docx", "png"),
    ("картинка.png", "jpg"), ("картинка.png", "webp"), ("картинка.png", "bmp"), ("картинка.png", "gif"),
    ("картинка.png", "tiff"), ("картинка.png", "ico"), ("картинка.png", "pdf"), ("фото.jpg", "png"),
    ("фото.jpg", "pdf"),
    ("звук.wav", "mp3"), ("звук.wav", "flac"), ("звук.wav", "ogg"), ("звук.wav", "m4a"),
    ("звук.wav", "opus"), ("звук.wav", "aac"),
    ("видео.mp4", "mp3"), ("видео.mp4", "webm"), ("видео.mp4", "mkv"), ("видео.mp4", "avi"),
    ("видео.mp4", "mov"), ("видео.mp4", "gif"),
]

ZIP_FORMATS = {"docx", "pptx", "xlsx", "odt", "ods", "odp", "epub"}
MAGIC = {
    "pdf": b"%PDF", "png": b"\x89PNG", "jpg": b"\xff\xd8", "gif": b"GIF8", "bmp": b"BM",
    "rtf": b"{\\rtf", "webp": b"RIFF", "ico": b"\x00\x00\x01\x00", "flac": b"fLaC", "ogg": b"OggS",
    "opus": b"OggS", "wav": b"RIFF", "mkv": b"\x1a\x45\xdf\xa3", "webm": b"\x1a\x45\xdf\xa3",
}


def check_format(path, ext):
    path = Path(path)
    assert path.exists() and path.stat().st_size > 0, f"пустой результат {path}"
    head = path.read_bytes()[:16]
    if ext in ZIP_FORMATS:
        assert zipfile.is_zipfile(path)
    elif ext in MAGIC:
        assert head.startswith(MAGIC[ext]), f"{path.name}: {head!r}"
    elif ext == "tiff":
        assert head[:2] in (b"II", b"MM")
    elif ext in ("mp4", "mov", "m4a"):
        assert head[4:8] == b"ftyp"
    elif ext == "avi":
        assert head.startswith(b"RIFF") and head[8:12] == b"AVI "
    elif ext in ("txt", "html", "md"):
        assert path.read_text(encoding="utf-8", errors="replace").strip()


@pytest.mark.parametrize("name,dst", CASES, ids=[f"{n}-{d}" for n, d in CASES])
def test_conversion(name, dst, data_dir, out_dir, engines):
    src = data_dir / name
    if not src.exists():
        pytest.skip("тестовый файл не создан (нет ffmpeg)")
    conv = registry.get(src.suffix, dst, engines)
    if conv is None:
        pytest.skip(f"нет доступного движка для {src.suffix} в {dst}")
    ctx = Context(namer=Namer(str(out_dir)))
    target = ctx.namer.file(src, dst)
    outputs = conv.run(engines, src, target, ctx)
    assert outputs
    for o in outputs:
        check_format(o, dst)


def test_pdf_text_layer(data_dir):
    from fileconverter.engines.pdf import pdf_info
    info = pdf_info(data_dir / "документ.pdf")
    assert info["pages"] == 2
    assert info["has_text"]


def test_slides_to_png_via_pdf(data_dir, out_dir, engines):
    if not engines.available("libreoffice"):
        pytest.skip("нет LibreOffice")
    from pptx import Presentation
    prs = Presentation()
    for t in ("Первый слайд", "Второй слайд"):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = t
    src = out_dir / "презентация.pptx"
    prs.save(str(src))
    conv = registry.get("pptx", "png", engines)
    assert conv.steps and conv.steps[0].dst == "pdf"
    ctx = Context(namer=Namer(str(out_dir)))
    outputs = conv.run(engines, src, ctx.namer.file(src, "png"), ctx)
    assert len(outputs) == 2
    for o in outputs:
        check_format(o, "png")
    pdf_conv = registry.get("pptx", "pdf", engines)
    out = pdf_conv.run(engines, src, ctx.namer.file(src, "pdf"), ctx)
    check_format(out[0], "pdf")


def test_xlsx_and_csv(out_dir, engines):
    if not engines.available("libreoffice"):
        pytest.skip("нет LibreOffice")
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Имя", "Сумма"])
    ws.append(["Иван", 100])
    src = out_dir / "таблица.xlsx"
    wb.save(str(src))
    ctx = Context(namer=Namer(str(out_dir)))
    for dst in ("pdf", "csv", "ods", "html"):
        conv = registry.get("xlsx", dst, engines)
        out = conv.run(engines, src, ctx.namer.file(src, dst), ctx)
        check_format(out[0], dst)
    csv = out_dir / "таблица.csv"
    assert "Иван" in csv.read_text(encoding="utf-8")


def test_heic_to_jpg(out_dir, engines):
    pillow_heif = pytest.importorskip("pillow_heif")
    from PIL import Image
    src = out_dir / "снимок.heic"
    img = Image.new("RGB", (64, 48), (200, 100, 50))
    try:
        heif = pillow_heif.from_pillow(img)
        heif.save(str(src), quality=80)
    except Exception as e:
        pytest.skip(f"кодировщик HEIC недоступен: {e}")
    conv = registry.get("heic", "jpg", engines)
    ctx = Context(namer=Namer(str(out_dir)))
    out = conv.run(engines, src, ctx.namer.file(src, "jpg"), ctx)
    check_format(out[0], "jpg")
