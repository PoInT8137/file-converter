import html
import logging
import re
import shutil
import threading
from functools import wraps
from pathlib import Path

import pymupdf

from ..config import assets_dir
from ..naming import pages_folder
from .base import ConvertError, ensure_output

PLEX_MONO = assets_dir() / "fonts" / "IBMPlexMono-Regular.ttf"

# PyMuPDF при импорте переводит MuPDF в однопоточный режим (reinit_singlethreaded),
# поэтому вся работа с ним из потоков задач и окна идёт по очереди
LOCK = threading.RLock()


def locked(func):
    @wraps(func)
    def wrapper(*args, **kwargs):
        with LOCK:
            return func(*args, **kwargs)
    return wrapper


def open_pdf(path, password=None):
    try:
        doc = pymupdf.open(str(path))
    except Exception as e:
        if "empty file" in str(e).lower():
            raise ConvertError("Файл пустой.", detail=str(e))
        raise ConvertError("Файл повреждён или это не PDF.", detail=str(e))
    if doc.needs_pass:
        if password and doc.authenticate(password):
            return doc
        doc.close()
        if password:
            raise ConvertError("Неверный пароль PDF")
        raise ConvertError("PDF защищён паролем. Сначала снимите пароль в разделе Инструменты PDF.")
    return doc


@locked
def pdf_info(path):
    """Число страниц, наличие текстового слоя и защита паролем."""
    info = {"pages": 0, "has_text": False, "encrypted": False}
    try:
        doc = pymupdf.open(str(path))
    except Exception:
        return info
    with doc:
        info["encrypted"] = bool(doc.needs_pass)
        if doc.needs_pass:
            return info
        info["pages"] = doc.page_count
        chars = 0
        for i in range(min(doc.page_count, 5)):
            chars += len(doc[i].get_text().strip())
            if chars > 30:
                break
        info["has_text"] = chars > 30
    return info


def save_doc(doc, dst, **kw):
    try:
        doc.save(str(dst), garbage=3, deflate=True, **kw)
    except Exception as e:
        raise ConvertError("Не удалось сохранить PDF.", detail=str(e))
    return ensure_output(dst)


class _PageLogHandler(logging.Handler):
    """pdf2docx пишет в лог строки вида (3/10) Page 3, по ним считаем прогресс."""

    pattern = re.compile(r"\((\d+)/(\d+)\)")

    def __init__(self, ctx):
        super().__init__()
        self.ctx = ctx
        self.stage = 0
        self.last = 0
        # лог общий, берём только записи своего потока
        self.thread = threading.get_ident()

    def emit(self, record):
        if record.thread != self.thread:
            return
        m = self.pattern.search(record.getMessage())
        if not m:
            return
        cur, total = int(m.group(1)), int(m.group(2))
        if cur < self.last:
            self.stage = 1
        self.last = cur
        try:
            self.ctx.progress(0.05 + 0.45 * self.stage + 0.45 * cur / max(total, 1))
        except ConvertError:
            # отмену из обработчика лога не пробросить, её проверит to_docx после конвертации
            pass


@locked
def to_docx(engines, src, dst, ctx):
    try:
        from pdf2docx import Converter
    except ImportError:
        raise ConvertError("Библиотека pdf2docx не установлена")
    doc = open_pdf(src)
    has_text = pdf_info(src)["has_text"]
    doc.close()
    if not has_text:
        ctx.log("Внимание: похоже на скан без текстового слоя, результат может быть пустым. Попробуйте OCR.")
    ctx.log("pdf2docx: разбор страниц, вёрстка может немного отличаться от оригинала")
    handler = _PageLogHandler(ctx)
    root = logging.getLogger()
    old_level = root.level
    root.addHandler(handler)
    root.setLevel(logging.INFO)
    try:
        cv = Converter(str(src))
        try:
            cv.convert(str(dst))
        finally:
            cv.close()
    except ConvertError:
        raise
    except Exception as e:
        raise ConvertError("pdf2docx не смог разобрать этот PDF. Попробуйте конвертацию в TXT или PNG.",
                           detail=str(e))
    finally:
        root.removeHandler(handler)
        root.setLevel(old_level)
    ctx.progress(1.0)
    return [ensure_output(dst)]


@locked
def to_txt(engines, src, dst, ctx):
    doc = open_pdf(src)
    parts = []
    with doc:
        for i, page in enumerate(doc):
            parts.append(page.get_text())
            ctx.progress((i + 1) / doc.page_count)
    text = "\n\n".join(p.strip() for p in parts)
    if not text.strip():
        ctx.log("Текст не найден: похоже на скан. Для таких файлов используйте распознавание (OCR).")
    Path(dst).write_text(text, encoding="utf-8")
    return [dst]


@locked
def to_html(engines, src, dst, ctx):
    doc = open_pdf(src)
    body = []
    with doc:
        for i, page in enumerate(doc):
            body.append(f'<div class="page" id="page{i + 1}">{page.get_text("xhtml")}</div>')
            ctx.progress((i + 1) / doc.page_count)
    title = html.escape(Path(src).stem)
    text = (
        "<!DOCTYPE html>\n<html><head><meta charset=\"utf-8\">"
        f"<title>{title}</title><style>body{{font-family:sans-serif;max-width:900px;margin:auto}}"
        ".page{border-bottom:1px solid #ccc;padding:20px 0}img{max-width:100%}</style></head><body>\n"
        + "\n".join(body) + "\n</body></html>\n"
    )
    Path(dst).write_text(text, encoding="utf-8")
    return [dst]


def render_pages(src, ctx, dpi=None):
    doc = open_pdf(src)
    dpi = dpi or ctx.dpi()
    with doc:
        total = doc.page_count
        for i, page in enumerate(doc):
            yield i, total, page.get_pixmap(dpi=dpi, alpha=False)


@locked
def to_images(engines, src, dst, ctx):
    dst = Path(dst)
    ext = dst.suffix.lower().lstrip(".")
    info = pdf_info(src)
    if info["encrypted"]:
        open_pdf(src)
    outputs = []
    folder = None
    for i, total, pix in render_pages(src, ctx):
        if total == 1:
            target = dst
        else:
            if folder is None:
                folder = pages_folder(dst, ctx.namer)
            target = folder / f"{dst.stem}-{i + 1:02d}.{ext}"
        if ext in ("jpg", "jpeg"):
            pix.save(str(target), jpg_quality=ctx.jpeg_quality())
        else:
            pix.save(str(target))
        outputs.append(target)
        ctx.progress((i + 1) / total)
    if folder:
        ctx.log(f"Страниц: {len(outputs)}, сохранены в папку {folder.name}")
    return outputs


@locked
def to_pptx(engines, src, dst, ctx):
    from pptx import Presentation
    from pptx.util import Emu

    ctx.log("Каждая страница станет картинкой на слайде, текст не будет редактироваться")
    doc = open_pdf(src)
    with doc:
        first = doc[0].rect
        prs = Presentation()
        prs.slide_width = Emu(int(first.width * 12700))
        prs.slide_height = Emu(int(first.height * 12700))
        blank = prs.slide_layouts[6]
        with ctx.tempdir() as tmp:
            for i, page in enumerate(doc):
                pix = page.get_pixmap(dpi=ctx.dpi(), alpha=False)
                img = Path(tmp) / f"p{i}.png"
                pix.save(str(img))
                slide = prs.slides.add_slide(blank)
                slide.shapes.add_picture(str(img), 0, 0, width=prs.slide_width, height=prs.slide_height)
                ctx.progress((i + 1) / doc.page_count * 0.95)
            prs.save(str(dst))
    ctx.progress(1.0)
    return [ensure_output(dst)]


@locked
def to_xlsx(engines, src, dst, ctx):
    from openpyxl import Workbook

    doc = open_pdf(src)
    wb = Workbook()
    wb.remove(wb.active)
    count = 0
    with doc:
        for i, page in enumerate(doc):
            try:
                tables = page.find_tables()
            except Exception:
                tables = []
            for t in tables:
                count += 1
                ws = wb.create_sheet(f"Стр{i + 1} табл{count}"[:31])
                for row in t.extract():
                    ws.append([("" if c is None else str(c)) for c in row])
            ctx.progress((i + 1) / doc.page_count)
    if count == 0:
        raise ConvertError("В PDF не найдено таблиц. Попробуйте конвертацию в DOCX или TXT.")
    wb.save(str(dst))
    ctx.log(f"Найдено таблиц: {count}")
    return [ensure_output(dst)]


@locked
def images_to_pdf(engines, files, dst, ctx):
    from PIL import Image

    from .images import open_image

    doc = pymupdf.open()
    with ctx.tempdir() as tmp:
        for i, f in enumerate(files):
            img = open_image(f)
            if img.mode not in ("RGB", "L"):
                img = flatten(img)
            tmpfile = Path(tmp) / f"{i}.jpg"
            img.save(tmpfile, quality=ctx.jpeg_quality())
            w, h = img.size
            page = doc.new_page(width=w * 72 / 96, height=h * 72 / 96)
            page.insert_image(page.rect, filename=str(tmpfile))
            ctx.progress((i + 1) / len(files) * 0.9)
        save_doc(doc, dst)
    doc.close()
    ctx.progress(1.0)
    return [dst]


def flatten(img):
    from PIL import Image

    img = img.convert("RGBA")
    bg = Image.new("RGB", img.size, "white")
    bg.paste(img, mask=img.split()[3])
    return bg


def image_to_pdf(engines, src, dst, ctx):
    return images_to_pdf(engines, [src], dst, ctx)


def parse_pages(text, total):
    """Разбирает строку вида '1-3, 5, 8-' в список индексов страниц (с нуля)."""
    text = (text or "").replace(" ", "")
    if not text:
        raise ConvertError("Не указаны страницы")
    result = []
    for part in re.split(r"[,;]", text):
        if not part:
            continue
        m = re.fullmatch(r"(\d*)-(\d*)", part)
        if m:
            a = int(m.group(1)) if m.group(1) else 1
            b = int(m.group(2)) if m.group(2) else total
        elif part.isdigit():
            a = b = int(part)
        else:
            raise ConvertError(f"Непонятный диапазон страниц: {part}")
        if a < 1 or b > total or a > b:
            raise ConvertError(f"Диапазон {part} выходит за пределы документа (страниц: {total})")
        result.extend(range(a - 1, b))
    if not result:
        raise ConvertError("Не указаны страницы")
    return result


@locked
def merge(engines, files, ctx, params):
    if len(files) < 2:
        raise ConvertError("Для объединения выберите хотя бы два PDF")
    dst = ctx.namer.file(files[0], "pdf", "объединённый")
    out = pymupdf.open()
    for i, f in enumerate(files):
        with open_pdf(f) as doc:
            out.insert_pdf(doc)
        ctx.progress((i + 1) / len(files) * 0.9)
    save_doc(out, dst)
    out.close()
    return [dst]


@locked
def split(engines, files, ctx, params):
    outputs = []
    for src in files:
        doc = open_pdf(src)
        with doc:
            ranges = (params.get("ranges") or "").strip()
            if ranges:
                groups = [parse_pages(g, doc.page_count) for g in ranges.split(";") if g.strip()]
            else:
                groups = [[i] for i in range(doc.page_count)]
            folder = ctx.namer.folder(src, "части")
            folder.mkdir(parents=True, exist_ok=True)
            for n, pages in enumerate(groups):
                part = pymupdf.open()
                for p in pages:
                    part.insert_pdf(doc, from_page=p, to_page=p)
                target = folder / f"{Path(src).stem}-{n + 1:02d}.pdf"
                save_doc(part, target)
                part.close()
                outputs.append(target)
                ctx.progress((n + 1) / len(groups))
        ctx.log(f"Частей: {len(groups)}, папка {folder.name}")
    return outputs


@locked
def extract_pages(engines, files, ctx, params):
    outputs = []
    for src in files:
        doc = open_pdf(src)
        with doc:
            pages = parse_pages(params.get("pages"), doc.page_count)
            out = pymupdf.open()
            for p in pages:
                out.insert_pdf(doc, from_page=p, to_page=p)
            dst = ctx.namer.file(src, "pdf", "страницы")
            save_doc(out, dst)
            out.close()
            outputs.append(dst)
    return outputs


@locked
def delete_pages(engines, files, ctx, params):
    outputs = []
    for src in files:
        doc = open_pdf(src)
        with doc:
            pages = set(parse_pages(params.get("pages"), doc.page_count))
            keep = [i for i in range(doc.page_count) if i not in pages]
            if not keep:
                raise ConvertError("Нельзя удалить все страницы документа")
            doc.select(keep)
            dst = ctx.namer.file(src, "pdf", "без страниц")
            save_doc(doc, dst)
            outputs.append(dst)
    return outputs


@locked
def rotate(engines, files, ctx, params):
    angle = int(params.get("angle", 90))
    outputs = []
    for src in files:
        doc = open_pdf(src)
        with doc:
            pages_text = (params.get("pages") or "").strip()
            pages = parse_pages(pages_text, doc.page_count) if pages_text else range(doc.page_count)
            for p in pages:
                page = doc[p]
                page.set_rotation((page.rotation + angle) % 360)
            dst = ctx.namer.file(src, "pdf", "повёрнутый")
            save_doc(doc, dst)
            outputs.append(dst)
    return outputs


@locked
def reorder(engines, files, ctx, params):
    outputs = []
    for src in files:
        doc = open_pdf(src)
        with doc:
            order = parse_pages(params.get("order"), doc.page_count)
            doc.select(order)
            dst = ctx.namer.file(src, "pdf", "порядок")
            save_doc(doc, dst)
            outputs.append(dst)
    return outputs


@locked
def compress(engines, files, ctx, params):
    level = params.get("level", "medium")
    settings = {"low": (150, 80), "medium": (120, 65), "high": (96, 50)}
    dpi, quality = settings.get(level, settings["medium"])
    outputs = []
    for src in files:
        before = Path(src).stat().st_size
        doc = open_pdf(src)
        with doc:
            try:
                doc.rewrite_images(dpi_threshold=dpi + 20, dpi_target=dpi, quality=quality)
            except Exception as e:
                ctx.log(f"Картинки не пересжаты: {e}")
            dst = ctx.namer.file(src, "pdf", "сжатый")
            try:
                doc.save(str(dst), garbage=4, deflate=True, deflate_images=True, deflate_fonts=True, clean=True)
            except Exception as e:
                raise ConvertError("Не удалось сохранить PDF.", detail=str(e))
        after = Path(dst).stat().st_size
        if after >= before:
            shutil.copy2(src, dst)
            ctx.log("Файл уже хорошо сжат, уменьшить размер не удалось")
        else:
            ctx.log(f"Размер: {human_size(before)} -> {human_size(after)}")
        outputs.append(dst)
    return outputs


@locked
def encrypt(engines, files, ctx, params):
    password = params.get("password") or ""
    if not password:
        raise ConvertError("Пароль не указан")
    outputs = []
    for src in files:
        doc = open_pdf(src)
        with doc:
            dst = ctx.namer.file(src, "pdf", "защищён")
            perm = pymupdf.PDF_PERM_PRINT | pymupdf.PDF_PERM_COPY | pymupdf.PDF_PERM_ACCESSIBILITY
            save_doc(doc, dst, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw=password,
                     owner_pw=password, permissions=perm)
            outputs.append(dst)
    return outputs


@locked
def decrypt(engines, files, ctx, params):
    password = params.get("password") or ""
    outputs = []
    for src in files:
        doc = open_pdf(src, password)
        with doc:
            dst = ctx.namer.file(src, "pdf", "без пароля")
            save_doc(doc, dst, encryption=pymupdf.PDF_ENCRYPT_NONE)
            outputs.append(dst)
    return outputs


@locked
def watermark(engines, files, ctx, params):
    text = (params.get("text") or "").strip()
    if not text:
        raise ConvertError("Текст водяного знака не указан")
    opacity = float(params.get("opacity", 0.25))
    font = pymupdf.Font(fontfile=str(PLEX_MONO)) if PLEX_MONO.exists() else pymupdf.Font("helv")
    outputs = []
    for src in files:
        doc = open_pdf(src)
        with doc:
            for i, page in enumerate(doc):
                r = page.rect
                size = max(18, min(r.width, r.height) / max(len(text), 4) * 1.3)
                width = font.text_length(text, fontsize=size)
                center = pymupdf.Point(r.width / 2, r.height / 2)
                tw = pymupdf.TextWriter(r, opacity=opacity, color=(0.5, 0.5, 0.5))
                tw.append(pymupdf.Point(center.x - width / 2, center.y + size / 3), text, font=font, fontsize=size)
                tw.write_text(page, morph=(center, pymupdf.Matrix(-35)))
                ctx.progress((i + 1) / doc.page_count)
            dst = ctx.namer.file(src, "pdf", "водяной знак")
            save_doc(doc, dst)
            outputs.append(dst)
    return outputs


@locked
def set_metadata(engines, files, ctx, params):
    outputs = []
    for src in files:
        doc = open_pdf(src)
        with doc:
            meta = dict(doc.metadata or {})
            for key in ("title", "author", "subject", "keywords"):
                if key in params:
                    meta[key] = params[key]
            doc.set_metadata(meta)
            dst = ctx.namer.file(src, "pdf", "метаданные")
            save_doc(doc, dst)
            outputs.append(dst)
    return outputs


@locked
def read_metadata(path):
    try:
        with pymupdf.open(str(path)) as doc:
            return dict(doc.metadata or {})
    except Exception:
        return {}


def extract_text(engines, files, ctx, params):
    outputs = []
    for src in files:
        dst = ctx.namer.file(src, "txt")
        outputs += to_txt(engines, src, dst, ctx)
    return outputs


@locked
def extract_images(engines, files, ctx, params):
    outputs = []
    for src in files:
        doc = open_pdf(src)
        folder = None
        seen = set()
        with doc:
            for i, page in enumerate(doc):
                for img in page.get_images(full=True):
                    xref = img[0]
                    if xref in seen:
                        continue
                    seen.add(xref)
                    data = doc.extract_image(xref)
                    if not data or not data.get("image"):
                        continue
                    if folder is None:
                        folder = ctx.namer.folder(src, "картинки")
                        folder.mkdir(parents=True, exist_ok=True)
                    target = folder / f"стр{i + 1:02d}-{len(seen):03d}.{data.get('ext', 'png')}"
                    target.write_bytes(data["image"])
                    outputs.append(target)
                ctx.progress((i + 1) / doc.page_count)
        if not outputs:
            raise ConvertError("В PDF нет встроенных картинок")
        ctx.log(f"Извлечено картинок: {len(outputs)}")
    return outputs


def human_size(n):
    for unit in ("Б", "КБ", "МБ", "ГБ"):
        if n < 1024 or unit == "ГБ":
            return f"{n:.0f} {unit}" if unit == "Б" else f"{n:.1f} {unit}"
        n /= 1024
